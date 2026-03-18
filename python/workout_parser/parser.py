"""Core parsing logic for workout Excel files."""

import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

EXERCISE_PATTERN = re.compile(r"^(\d+)x(\d+)(?:x(\d+))?$")
DAY_PATTERN = re.compile(r"Day\s+(\d+)")
WEEK_PATTERN = re.compile(r"Week\s+(\d+)")
FILENAME_PATTERN = re.compile(r"^([A-Za-z]+)-(\d{4})-")


@dataclass
class Exercise:
    date: date
    order: str
    name: str
    sets: int
    reps: int
    weight: int


def first_monday(year: int, month: int) -> date:
    """Return the first Monday of the given month."""
    d = date(year, month, 1)
    days_ahead = (7 - d.weekday()) % 7  # Monday is 0
    return d + timedelta(days=days_ahead)


def parse_month_year(filename: str) -> tuple[int, int]:
    """Extract month and year from a filename like 'December-2024-4-Day-...'."""
    match = FILENAME_PATTERN.match(filename)
    if not match:
        raise ValueError(f"Cannot parse month/year from filename: {filename}")
    month_name, year_str = match.group(1), match.group(2)
    month_num = datetime.strptime(month_name, "%B").month
    return int(year_str), month_num


def parse_workbook(path: Path) -> list[Exercise]:
    """Parse a single workout Excel file and return a list of exercises."""
    year, month = parse_month_year(path.name)
    monday = first_monday(year, month)
    exercises: list[Exercise] = []

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet_name in wb.sheetnames:
            week_match = WEEK_PATTERN.match(sheet_name)
            if not week_match:
                continue

            week_num = int(week_match.group(1))
            week_date = monday + timedelta(weeks=week_num - 1)
            current_date = week_date

            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=1):
                exercise_cell = row[1].value if len(row) > 1 else None
                set_cell = row[5].value if len(row) > 5 else None

                if not exercise_cell:
                    continue

                exercise_val = str(exercise_cell).strip()

                day_match = DAY_PATTERN.match(exercise_val)
                if day_match:
                    current_date = week_date + timedelta(
                        days=int(day_match.group(1))
                    )

                if not set_cell:
                    continue

                set_val = str(set_cell).strip()
                set_match = EXERCISE_PATTERN.match(set_val)
                if set_match:
                    if ":" not in exercise_val:
                        continue
                    order, name = exercise_val.split(":", 1)
                    exercises.append(
                        Exercise(
                            date=current_date,
                            order=order.strip(),
                            name=name.strip(),
                            sets=int(set_match.group(1)),
                            reps=int(set_match.group(2)),
                            weight=int(set_match.group(3) or 0),
                        )
                    )
    finally:
        wb.close()

    return exercises


def parse_folder(folder: Path) -> pd.DataFrame:
    """Parse all workout Excel files in a folder and return a consolidated DataFrame."""
    all_exercises: list[Exercise] = []
    for path in sorted(folder.glob("*.xlsx")):
        all_exercises.extend(parse_workbook(path))

    if not all_exercises:
        return pd.DataFrame(columns=["Date", "Order", "Name", "Sets", "Reps", "Weight"])

    df = pd.DataFrame(
        [
            {
                "Date": e.date,
                "Order": e.order,
                "Name": e.name,
                "Sets": e.sets,
                "Reps": e.reps,
                "Weight": e.weight,
            }
            for e in all_exercises
        ]
    )
    return df.sort_values("Date").reset_index(drop=True)


def write_output(df: pd.DataFrame, output_path: Path) -> None:
    """Write the consolidated DataFrame to an Excel file."""
    df.to_excel(output_path, index=False, sheet_name="Monthly Exercises")